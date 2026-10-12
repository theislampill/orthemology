from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
from datetime import datetime,timezone
import json,re
from docx import Document
ROOT=Path('private-locator:8f24556a330b05a11a95');OUT=ROOT/'t20-continuation-20261008/owner-review'
base=json.loads((ROOT/'recovery-t20/checkpoint/Orthemology_T20_Continuation_Owner_Review_Checkpoint_20261008/tranche20/research-frontier-continuation/closure-case-current/OBLIGATION_MATRIX.json').read_text())
now=json.loads((OUT/'REVIEW_OBLIGATIONS.json').read_text())
for key,idfield in [('original_obligations','id'),('formal_objective_submatrix','id'),('twelve_avenue_preservation','number')]:
 assert {r[idfield]:r['status'] for r in base[key]}=={r[idfield]:r['status'] for r in now[key]},key
assert len(now['original_obligations'])==20
assert len(now['formal_objective_submatrix'])==11
assert len(now['twelve_avenue_preservation'])==12
p=OUT/'Orthemology_T20_Updated_Owner_Review_20261008.docx'
d=Document(p);t='\n'.join(x.text for x in d.paragraphs)
for r in now['original_obligations']:assert re.search(r'\b'+r['id']+r'\s+'+r['status'].upper().replace('_',' '),t)
for s in ['T20 remains open','Nineteen scoped effect-contract checks','sixteen scoped checks','UNVERIFIED','concentration','three and four occurrences','nineteen checks','eighteen-check receipt','nondefault event states','ledger’s own issued-act profile']:
 assert s in t,s
assert 'FINAL_EFFECT' not in t
assert 'authorizes' not in t and 'normalization' not in t
assert not any(re.search(r'[^A-Za-z0-9 ]',p.text) for p in d.paragraphs if p.style.name in ['Title','Heading 1','Heading 2'])
with ZipFile(p) as z:
 assert b'w:pBdr' not in z.read('word/document.xml')
 assert b'w:pBdr' not in z.read('word/styles.xml')
 assert b'w:documentProtection' not in z.read('word/settings.xml')
 assert b'w:val="en-GB"' in z.read('word/styles.xml')
files=['Orthemology_T20_Updated_Owner_Review_20261008.docx','Orthemology_T20_Updated_Owner_Review_Source_20261008.md','REVIEW_OBLIGATIONS.json']
manifest={'schema':'t20-owner-review-delivery/1','date_utc':datetime.now(timezone.utc).isoformat(),'status':'rendered_and_content_verified_awaiting_visual_signoff','original_status_values_preserved':True,'original_obligations':20,'formal_rows':11,'avenues':12,'expected_render_pages':9,'research_status':'T20_OPEN_OWNER_ACCEPTANCE_PENDING','historical_six_hour_floor':'UNVERIFIED','no_final_integration':True,'outputs':[{'path':f,'bytes':(OUT/f).stat().st_size,'sha256':sha256((OUT/f).read_bytes()).hexdigest()} for f in files],'exclude_from_user_delivery':['*_DRAFT.json','REVIEW_CONTENT.json','prepare_review.py','finalise_review.py','build_review.py','verify_review.py','render/','render-final/','render-19/','render-release/']}
(OUT/'OWNER_REVIEW_DELIVERY.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps(manifest,indent=2))
