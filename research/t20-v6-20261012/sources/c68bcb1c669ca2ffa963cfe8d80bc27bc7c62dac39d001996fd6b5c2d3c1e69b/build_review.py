from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
OUT=Path('private-locator:a12b7fb1f25ab7f5d347')
content=json.loads((OUT/'REVIEW_CONTENT.json').read_text())
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
styles=d.styles
for n in ['Normal','Title','Heading 1','Heading 2','Subtitle']:
 styles[n].font.name='Calibri';styles[n].font.color.rgb=RGBColor.from_string('111111')
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Normal'].paragraph_format.line_spacing=1.04
styles['Title'].font.size=Pt(23)
styles['Title'].paragraph_format.space_after=Pt(9)
styles['Heading 1'].font.size=Pt(18)
styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Heading 1'].paragraph_format.space_after=Pt(10)
styles['Heading 2'].font.size=Pt(11)
styles['Heading 2'].paragraph_format.space_before=Pt(7)
styles['Heading 2'].paragraph_format.space_after=Pt(3)
for n in ['Heading 1','Heading 2']:
 styles[n].paragraph_format.keep_with_next=True
styles['Normal'].paragraph_format.widow_control=True
# Match the reference's restrained hierarchy, with no title rules or tables.
for i,pg in enumerate(content):
 title=d.add_paragraph(pg['title'],style='Title' if i==0 else 'Heading 1')
 if i:title.paragraph_format.page_break_before=True
 blocks=pg['blocks']
 # Compact accounting rows into one continuous paragraph each, preserving every word.
 if i in (3,4):
  merged=[];j=0
  while j<len(blocks):
   b=blocks[j]
   if b['kind']=='sub' and j+1<len(blocks) and blocks[j+1]['kind']=='p':
    merged.append({'kind':'inline','title':b['title'],'text':b['text']+' '+blocks[j+1]['text']});j+=2
   else:merged.append(b);j+=1
  blocks=merged
 for b in blocks:
  if b['kind']=='inline':
   p0=d.add_paragraph();r=p0.add_run(b['title']+'. ');r.bold=True;p0.add_run(b['text'])
   p0.paragraph_format.space_after=Pt(9)
  elif b['kind']=='p':
   para=d.add_paragraph(b['text'])
   if i==0 and b==pg['blocks'][0]:
    para.runs[0].italic=True;para.runs[0].font.size=Pt(10)
  elif b['kind']=='sub':
   d.add_paragraph(b['title'].replace('-',' ').replace('–',' '),style='Heading 2')
   d.add_paragraph(b['text'])
  elif b['kind']=='links':
   p=d.add_paragraph()
   for j,(label,url) in enumerate(b['items']):
    if j:p.add_run('  |  ')
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'245C86');pr.append(c)
    s=OxmlElement('w:sz');s.set(qn('w:val'),'18');pr.append(s);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
 # Dense accounting pages use consistent 10 point body text, not compressed columns.
 if i in (3,4,8):
  # The range to restyle is identified by the most recent page heading.
  idx=max(k for k,p0 in enumerate(d.paragraphs) if p0.text==pg['title'])
  for p0 in d.paragraphs[idx+1:]:
   if p0.style.name=='Normal':
    p0.paragraph_format.space_after=Pt(4)
    for r in p0.runs:r.font.size=Pt(10)
# A page field is useful in a nine-page review case.
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Orthemology T20  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.title='Orthemology T20 owner review'
d.core_properties.subject='Renewed investigation and original obligation accounting'
d.core_properties.author=''
d.core_properties.keywords='T20; owner review; open; 8 October 2026'
# Remove Word's built-in decorative title border and any direct paragraph rules.
for elem in list(d.styles.element.iter(qn('w:pBdr'))) + list(d.element.iter(qn('w:pBdr'))):
 elem.getparent().remove(elem)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-GB');styles['Normal'].element.get_or_add_rPr().append(lang)
p=OUT/'Orthemology_T20_Updated_Owner_Review_20261008.docx';d.save(p)
print(p)
