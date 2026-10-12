from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
styles=d.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
 st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 if name=='Normal':st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
 elif name=='Title':st.font.size=Pt(23);st.paragraph_format.space_after=Pt(10)
 elif name=='Heading 1':st.font.size=Pt(14);st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
 elif name=='Heading 2':st.font.size=Pt(12);st.paragraph_format.keep_with_next=True
 else:st.font.size=Pt(10.5)
# Clean black title styling with no inherited paragraph rule.
for style in [styles['Title'],styles['Subtitle']]:
 ppr=style.element.find(qn('w:pPr'))
 if ppr is not None:
  for el in list(ppr):
   if el.tag==qn('w:pBdr'):ppr.remove(el)
text=(R/'AUDIT_REPORT.md').read_text();buf=[]
def flush():
 if buf:
  p=d.add_paragraph(' '.join(buf));p.paragraph_format.widow_control=True;buf.clear()
for line in text.splitlines():
 if line.startswith('# '):
  flush();d.add_paragraph(line[2:],style='Title');d.add_paragraph('Research shadows and qualified mathematical extensions\n30 September 2026',style='Subtitle')
 elif line.startswith('## '):flush();d.add_paragraph(line[3:],style='Heading 1')
 elif line.startswith('- '):
  flush();p=d.add_paragraph(line[2:],style='List Bullet');p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.05
 elif not line.strip():flush()
 else:buf.append(line)
flush()
p=sec.footer.paragraphs[0];p.alignment=2
r=p.add_run('Orthemology formal qualification audit  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
d.core_properties.title='Orthemology formal qualification audit';d.core_properties.author='dot';d.core_properties.subject='Independent bounded technical review of repaired research shadows'
d.save(R/'Orthemology_Formal_Qualification_Audit.docx')
print(R/'Orthemology_Formal_Qualification_Audit.docx')
